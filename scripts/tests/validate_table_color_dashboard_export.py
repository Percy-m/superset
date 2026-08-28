# Licensed to the Apache Software Foundation (ASF) under one
# or more contributor license agreements.  See the NOTICE file
# distributed with this work for additional information
# regarding copyright ownership.  The ASF licenses this file
# to you under the Apache License, Version 2.0 (the
# "License"); you may not use this file except in compliance
# with the License.  You may obtain a copy of the License at
#
#   http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing,
# software distributed under the License is distributed on an
# "AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY
# KIND, either express or implied.  See the License for the
# specific language governing permissions and limitations
# under the License.
"""Verify real local Dashboard XLSX exports against a displayed color snapshot.

Run after the local Superset service has loaded the implementation::

    SUPERSET_CONFIG_PATH=$PWD/superset_config.py venv/bin/python \
        scripts/tests/validate_table_color_dashboard_export.py

Uses the existing local FR-03 dashboard and saved charts without editing them.
Only result caches and the export endpoint's temporary files are created.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import logging
import os
import sys
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from openpyxl import load_workbook


def main() -> int:
    """Exercise the running HTTP API, never mocked SQL or a modified dashboard."""
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
    from validate_table_color_filters import (  # pylint: disable=import-outside-toplevel
        LocalApi,
        require,
    )

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-url", default="http://127.0.0.1:8088")
    parser.add_argument("--dashboard-id", type=int, default=3)
    parser.add_argument("--chart-id", type=int, default=3)
    parser.add_argument("--tab-id", default="TAB-FR03-OUTER")
    parser.add_argument("--username", default="admin")
    args = parser.parse_args()
    endpoint = urlparse(args.base_url)
    require(
        endpoint.scheme in {"http", "https"}
        and endpoint.hostname in {"localhost", "127.0.0.1", "::1"}
        and endpoint.username is None,
        "only loopback service URLs are permitted",
    )
    require(os.environ.get("SUPERSET_CONFIG_PATH"), "local configuration is required")
    for name in ("NO_PROXY", "no_proxy"):
        os.environ[name] = ",".join(
            dict.fromkeys(
                [*os.environ.get(name, "").split(","), "127.0.0.1", "localhost", "::1"]
            )
        )
    logging.disable(logging.CRITICAL)
    from superset.app import create_app  # pylint: disable=import-outside-toplevel

    app = create_app()
    from superset.extensions import (  # pylint: disable=import-outside-toplevel
        db,
        security_manager,
    )
    from superset.models.dashboard import (  # pylint: disable=import-outside-toplevel
        Dashboard,
    )

    with app.app_context():
        dashboard = db.session.query(Dashboard).filter_by(id=args.dashboard_id).one()
        chart = next(
            (item for item in dashboard.slices if item.id == args.chart_id), None
        )
        require(
            chart is not None and chart.viz_type == "table", "saved Table is absent"
        )
        assert chart is not None
        require(args.tab_id in dashboard.position, "selected Tab is absent")
        require(
            all(
                item.datasource.database.url_object.host
                in {"localhost", "127.0.0.1", "::1"}
                for item in dashboard.slices
            ),
            "dashboard datasets must be local",
        )
        before = hashlib.sha256(
            json.dumps([dashboard.position_json, chart.params]).encode()
        ).hexdigest()
        user = security_manager.find_user(username=args.username)
        require(
            user is not None and user.is_active, "local test account is unavailable"
        )
        api = LocalApi(app, user.id, args.base_url, "http")
        payload = json.loads(chart.query_context)
        payload.update(
            form_data={
                **chart.form_data,
                "dashboardId": dashboard.id,
                "extra_form_data": {},
            },
            result_type="full",
            result_format="json",
            force=True,
        )
        for query in payload["queries"]:
            query.pop("alert_filters", None)
            query.pop("is_table_alert_totals", None)
            query.pop("table_color_filter", None)
        main_query = payload["queries"][0]
        main_query["row_offset"] = 0
        main_query["row_limit"] = 20
        main_query["table_color_filter"] = {"version": 2, "selections": []}
        original = api.request("POST", "/api/v1/chart/data", payload)
        require(original.status == 200, "dashboard snapshot request failed")
        primary = original.body["result"][0]
        metadata = primary["table_color_metadata"]
        require(metadata["status"] == "ready", "dashboard snapshot unavailable")
        selections = [{"column": "gross_revenue", "colors": ["GREEN"]}]
        selected = copy.deepcopy(payload)
        selected["force"] = False
        selected["queries"][0]["table_color_filter"].update(
            selections=selections, snapshot_id=metadata["snapshot_id"]
        )
        selected_response = api.request("POST", "/api/v1/chart/data", selected)
        require(selected_response.status == 200, "dashboard color selection failed")
        selected_metadata = selected_response.body["result"][0]["table_color_metadata"]
        require(
            selected_metadata["filtered_rowcount"] > 0, "fixture has no matching rows"
        )
        expected_query = copy.deepcopy(selected)
        expected_query["result_type"] = "results"
        expected = api.request("POST", "/api/v1/chart/data", expected_query)
        require(expected.status == 200, "snapshot export reference failed")
        expected_primary = expected.body["result"][0]
        export_payload: dict[str, Any] = {
            "tabIds": [args.tab_id],
            "dataMask": {
                str(chart.id): {
                    "ownState": {
                        "alertFilter": {"version": 2, "selections": selections}
                    }
                }
            },
            "colorSnapshots": {
                str(chart.id): {
                    "snapshot_id": metadata["snapshot_id"],
                    "generation": metadata["generation"],
                    "theme_mode": metadata.get("theme_mode", "default"),
                }
            },
        }
        route = f"/api/v1/dashboard/{dashboard.id}/export_xlsx/"
        response = api.request("POST", route, export_payload)
        require(response.status == 200, f"Tab XLSX returned HTTP {response.status}")
        require(response.content.startswith(b"PK"), "Tab response is not a workbook")
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        sheet = workbook.worksheets[0]
        expected_rows = expected_primary["data"]
        columns = expected_primary["colnames"]
        require(
            [cell.value for cell in sheet[1]] == columns,
            "Tab columns do not match the displayed snapshot",
        )
        actual_rows = list(
            sheet.iter_rows(min_row=2, max_row=1 + len(expected_rows), values_only=True)
        )
        require(
            actual_rows
            == [tuple(row.get(column) for column in columns) for row in expected_rows],
            "Tab rows or order differ from the selected snapshot",
        )
        color_index = columns.index("gross_revenue") + 1
        for row_index, paint in enumerate(
            expected_primary["table_color_metadata"]["styles"], start=2
        ):
            expected_color = (
                paint["gross_revenue"]["backgroundColor"].lstrip("#")[:6].upper()
            )
            require(
                sheet.cell(row_index, color_index).fill.fgColor.rgb[-6:]
                == expected_color,
                "Tab painting differs from the displayed snapshot",
            )
        sheet_names = list(workbook.sheetnames)
        workbook.close()
        require(len(sheet_names) >= 2, "nested Tab did not export multiple sheets")
        print(
            json.dumps(
                {
                    "case": "DASHBOARD_NESTED_COLOR_XLSX",
                    "status": "PASS",
                    "selected_rows": len(expected_rows),
                    "sheets": len(sheet_names),
                }
            )
        )
        # A broken snapshot must yield JSON, never a successful partial workbook.
        stale = copy.deepcopy(export_payload)
        stale["colorSnapshots"][str(chart.id)]["snapshot_id"] = (
            "expired-color-acceptance-reference"
        )
        failed = api.request("POST", route, stale)
        require(failed.status == 410, "expired Tab snapshot did not fail explicitly")
        require(
            failed.body.get("error_code") == "TABLE_COLOR_FILTER_SNAPSHOT_EXPIRED",
            "incorrect Tab expiry error",
        )
        require(
            not failed.content.startswith(b"PK"), "partial XLSX escaped after failure"
        )
        print(
            json.dumps(
                {
                    "case": "DASHBOARD_ATOMIC_FAILURE",
                    "status": "PASS",
                    "status_code": failed.status,
                }
            )
        )
        db.session.refresh(dashboard)
        db.session.refresh(chart)
        after = hashlib.sha256(
            json.dumps([dashboard.position_json, chart.params]).encode()
        ).hexdigest()
        require(before == after, "saved dashboard or formatter was modified")
        print(
            json.dumps(
                {
                    "summary": "PASS",
                    "api_requests": api.requests,
                    "saved_objects_changed": 0,
                }
            )
        )
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        # Keep arbitrary API/SQL exception bodies out of acceptance output.
        print(json.dumps({"summary": "FAIL", "error_type": type(error).__name__}))
        raise SystemExit(1) from None
