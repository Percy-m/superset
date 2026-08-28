# Licensed to the Apache Software Foundation (ASF) under one
# or more contributor license agreements.  See the NOTICE file
# distributed with this work for additional information
# regarding copyright ownership.  The ASF licenses this file
# to you under the Apache License, Version 2.0 (the
# "License"); you may not use this file except in compliance
# with the License.  You may obtain a copy of the License at
#
#   http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing,
# software distributed under the License is distributed on an
# "AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY
# KIND, either express or implied.  See the License for the
# specific language governing permissions and limitations
# under the License.
"""Exercise real local ClickHouse color-filter APIs without changing saved charts.

Run from the repository root using its configured development environment::

    SUPERSET_CONFIG_PATH=$PWD/superset_config.py venv/bin/python \
        scripts/tests/validate_table_color_filters.py --transport http
    SUPERSET_CONFIG_PATH=$PWD/superset_config.py venv/bin/python \
        scripts/tests/validate_table_color_filters.py --transport test-client
    SUPERSET_CONFIG_PATH=$PWD/superset_config.py venv/bin/python \
        scripts/tests/validate_table_color_filters.py --transport http --cell-bars-only

The HTTP mode checks the running service. The test-client mode runs the same
real Flask routes, ClickHouse queries and Redis cache, with a pass-through
engine execute observer counting actual SELECTs; nothing is stubbed. Both
modes create/delete only owned temporary Explore drafts. Existing Slice and
Dataset configuration and fixture rows are never edited. Result snapshots
expire under their configured TTL. Output contains case IDs and counts, never
tokens, SQL, source rows or filter values.

The fixture suite uses the saved sales Table (chart 3 by default) and the saved
customer Table (chart 1). Cross-month tests use the customer's declared
``signup_date`` datetime column, not the sales fixture's unmarked date fields.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import logging
import math
import os
import sys
import time
import uuid
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import timedelta
from io import BytesIO
from pathlib import Path
from typing import Any, Iterator, TYPE_CHECKING
from urllib.parse import urlparse

import requests
from flask_jwt_extended import create_access_token
from openpyxl import load_workbook
from PIL import ImageColor

if TYPE_CHECKING:
    from flask import Flask
    from flask.testing import FlaskClient
    from openpyxl.cell.cell import Cell
    from openpyxl.formatting.rule import DataBar


class AcceptanceError(Exception):
    """A failure containing only a non-sensitive, predetermined diagnostic."""


def require(condition: object, message: str) -> None:
    """Fail without printing API response bodies or source records."""
    if not condition:
        raise AcceptanceError(message)


@dataclass
class ApiResponse:
    """A small transport-neutral HTTP response view."""

    status: int
    body: dict[str, Any]
    content: bytes
    headers: dict[str, str]


@dataclass(frozen=True)
class SourceChart:
    """Retain source config without a request-scoped SQLAlchemy object."""

    id: int
    datasource_id: int
    form_data: dict[str, Any]
    params: str
    query_context: dict[str, Any]
    engine_spec: Any


class LocalApi:
    """Authenticate using a short-lived token minted by the local app config."""

    def __init__(self, app: Flask, user_id: int, base_url: str, transport: str) -> None:
        self.app = app
        self.user_id = user_id
        self.base_url = base_url.rstrip("/")
        self.transport = transport
        self.session = requests.Session()
        self.session.trust_env = False
        self.client: FlaskClient = app.test_client()
        token = create_access_token(
            identity=str(user_id), expires_delta=timedelta(minutes=15)
        )
        self.headers = {"Authorization": f"Bearer {token}"}
        self.requests = 0
        csrf = self.request("GET", "/api/v1/security/csrf_token/")
        require(csrf.status == 200, "local token/CSRF authentication failed")
        self.headers["X-CSRFToken"] = csrf.body["result"]

    def request(
        self,
        method: str,
        path: str,
        payload: dict[str, Any] | None = None,
        *,
        anonymous: bool = False,
    ) -> ApiResponse:
        """Perform one request without exposing its body or auth headers."""
        self.requests += 1
        headers = {} if anonymous else self.headers
        if self.transport == "http":
            response = self.session.request(
                method,
                self.base_url + path,
                json=payload,
                headers=headers,
                timeout=60,
                allow_redirects=False,
            )
            content = response.content
            try:
                parsed = response.json()
            except ValueError:
                parsed = {}
            return ApiResponse(
                response.status_code,
                parsed if isinstance(parsed, dict) else {},
                content,
                dict(response.headers),
            )
        # Each HTTP request needs a fresh g/JWT context, just like the server.
        with self.app.app_context():
            local_response = self.client.open(
                path, method=method, json=payload, headers=headers
            )
        return ApiResponse(
            local_response.status_code,
            local_response.get_json(silent=True) or {},
            local_response.data,
            dict(local_response.headers),
        )


class FixtureQueryCounter:
    """Observe actual ClickHouse fixture SELECTs without retaining SQL text."""

    def __init__(self) -> None:
        self.count = 0

    @contextmanager
    def observe(self, engine_spec: Any) -> Iterator[None]:
        """Observe raw-cursor execution while delegating every real database call."""
        original = engine_spec.execute
        descriptor = engine_spec.__dict__.get("execute")

        def execute(
            cls: Any, cursor: Any, query: str, database: Any, **kwargs: Any
        ) -> None:
            if query.lstrip().upper().startswith("SELECT") and any(
                name in query.lower()
                for name in ("fact_sales", "dim_customer", "fact_events")
            ):
                self.count += 1
            original(cursor, query, database, **kwargs)

        engine_spec.execute = classmethod(execute)
        try:
            yield
        finally:
            if descriptor is None:
                delattr(engine_spec, "execute")
            else:
                engine_spec.execute = descriptor


class ColorAcceptance:
    """Check selection against independent expected rows and frozen baseline paint."""

    def __init__(self, api: LocalApi, chart: SourceChart) -> None:
        self.api = api
        self.chart = chart
        self.run_id = uuid.uuid4().hex
        self.drafts: list[str] = []
        self.counter = FixtureQueryCounter()
        self.outcomes: list[dict[str, Any]] = []
        self.saved_fingerprint = self._fingerprint()
        self.protected_charts = {chart.id: self.saved_fingerprint}

    def _fingerprint(self) -> str:
        return hashlib.sha256((self.chart.params or "").encode("utf-8")).hexdigest()

    def record(self, case_id: str, **counts: object) -> None:
        """Emit only controlled case labels, transport and numerical diagnostics."""
        outcome = {"case": case_id, "status": "PASS", **counts}
        self.outcomes.append(outcome)
        print(json.dumps(outcome, ensure_ascii=False), flush=True)

    @staticmethod
    def rule(
        column: str,
        scheme: str,
        operator: str,
        target: object = 0,
        **options: object,
    ) -> dict[str, Any]:
        """Build native formatter controls without retired subject/level fields."""
        return {
            "column": column,
            "colorScheme": scheme,
            "operator": operator,
            "targetValue": target,
            "useGradient": False,
            "filterable": True,
            **options,
        }

    def raw_rules(self) -> list[dict[str, Any]]:
        """Cover three palettes, strings and a second numerical target."""
        return [
            self.rule("revenue", "colorSuccess", ">", 100),
            self.rule(
                "revenue",
                "colorWarning",
                "≤ x ≤",
                targetValueLeft=0,
                targetValueRight=100,
            ),
            self.rule("revenue", "colorError", "<", 0),
            self.rule("region", "colorWarning", "begins with", "A"),
            self.rule("quantity", "colorSuccess", ">", 4),
        ]

    def create_draft(
        self, form_data: dict[str, Any], datasource_id: int, chart_id: int
    ) -> str:
        """Store only owned temporary state and register it for cleanup."""
        created = self.api.request(
            "POST",
            "/api/v1/explore/form_data",
            {
                "datasource_id": datasource_id,
                "datasource_type": "table",
                "chart_id": chart_id,
                "form_data": json.dumps(form_data),
            },
        )
        require(created.status == 201, f"draft creation returned HTTP {created.status}")
        key = created.body["key"]
        self.drafts.append(key)
        return key

    def make_payload(
        self,
        *,
        server: bool,
        rows: int = 180,
        aggregate: bool = False,
        percentage: bool = False,
        float_percentage: bool = False,
        show_cell_bars: bool = False,
        rules: list[dict[str, Any]] | None = None,
    ) -> dict[str, Any]:
        """Create an owned draft and a matching native Chart Data request."""
        columns = (
            ["region", "channel", "alert_bucket", "quantity"]
            if aggregate
            else ["sale_id", "region", "quantity", "revenue"]
        )
        metrics: list[object] = (
            [
                "gross_revenue",
                {
                    "expressionType": "SQL",
                    "sqlExpression": "SUM(revenue - cost)",
                    "label": "qa_profit",
                    "hasCustomLabel": True,
                },
            ]
            if aggregate
            else []
        )
        percent_metric = "qa_float_revenue" if float_percentage else "gross_revenue"
        if float_percentage:
            metrics.append(
                {
                    "expressionType": "SQL",
                    "sqlExpression": "SUM(toFloat64(revenue))",
                    "label": percent_metric,
                    "hasCustomLabel": True,
                }
            )
        formatting = rules if rules is not None else self.raw_rules()
        form_data = {
            **copy.deepcopy(self.chart.form_data),
            "viz_type": "table",
            "slice_id": self.chart.id,
            "datasource": f"{self.chart.datasource_id}__table",
            "query_mode": "aggregate" if aggregate else "raw",
            "server_pagination": server,
            "server_page_length": 20,
            "row_limit": 50_000,
            "all_columns": [] if aggregate else columns,
            "groupby": columns if aggregate else [],
            "metrics": metrics,
            "percent_metrics": [percent_metric] if percentage else [],
            "percent_metric_calculation": "all_records" if percentage else "row_limit",
            "show_totals": percentage,
            "show_cell_bars": show_cell_bars,
            "time_range": "No filter",
            "time_compare": [],
            "comparison_type": "values",
            "granularity_sqla": "sale_date",
            "adhoc_filters": [],
            "column_config": {},
            "conditional_formatting": formatting,
            "_color_acceptance_run": self.run_id,
            "_color_acceptance_case": uuid.uuid4().hex,
        }
        if show_cell_bars:
            form_data.update(align_pn=True, color_pn=True)
        key = self.create_draft(form_data, self.chart.datasource_id, self.chart.id)
        filters = [{"col": "sale_id", "op": "<=", "val": rows}]
        query = {
            "columns": columns,
            "metrics": metrics,
            "time_range": "No filter",
            "granularity": "sale_date",
            "time_offsets": [],
            "filters": filters,
            "orderby": [[column, True] for column in columns],
            "row_limit": 20 if server else 50_000,
            "row_offset": 0,
            "post_processing": (
                [
                    {
                        "operation": "contribution",
                        "options": {
                            "columns": [percent_metric],
                            "rename_columns": [f"%{percent_metric}"],
                        },
                    }
                ]
                if percentage
                else []
            ),
        }
        queries = [query]
        if server:
            queries.append(
                {
                    **copy.deepcopy(query),
                    "is_rowcount": True,
                    "time_offsets": [],
                    "post_processing": [],
                    "row_limit": 50_000,
                }
            )
        if percentage:
            queries.append(
                {
                    **copy.deepcopy(query),
                    "columns": [],
                    "orderby": [],
                    "post_processing": [],
                    "row_limit": 50_000,
                }
            )
        queries[0]["table_color_filter"] = {
            "version": 2,
            "form_data_key": key,
            "selections": [],
        }
        return {
            "datasource": {"id": self.chart.datasource_id, "type": "table"},
            "force": True,
            "form_data": form_data,
            "queries": queries,
            "result_format": "json",
            "result_type": "full",
        }

    def query(
        self, payload: dict[str, Any], status: int = 200, code: str | None = None
    ) -> ApiResponse:
        """Require the expected HTTP outcome without displaying backend errors."""
        response = self.api.request("POST", "/api/v1/chart/data", payload)
        actual_code = response.body.get("error_code", "NO_CODE")
        require(
            response.status == status,
            f"Chart Data returned HTTP {response.status}, expected {status}; "
            f"code={actual_code}",
        )
        if code:
            require(actual_code == code, "unexpected structured error code")
        return response

    @staticmethod
    def primary(response: ApiResponse) -> dict[str, Any]:
        """Require a usable Chart Data result and versioned color metadata."""
        results = response.body.get("result")
        require(isinstance(results, list) and results, "missing Chart Data results")
        assert isinstance(results, list)
        primary = results[0]
        require(primary.get("status") == "success", "Chart Data query did not succeed")
        require("table_color_metadata" in primary, "color metadata missing")
        return primary

    @staticmethod
    def request_for(
        payload: dict[str, Any],
        snapshot_id: str,
        selections: list[dict[str, Any]],
        *,
        offset: int = 0,
    ) -> dict[str, Any]:
        """Change only snapshot selection and presentation offset."""
        selected = copy.deepcopy(payload)
        selected["force"] = False
        selected["queries"][0]["row_offset"] = offset
        selected["queries"][0]["table_color_filter"].update(
            snapshot_id=snapshot_id, selections=selections
        )
        return selected

    def complete(
        self, payload: dict[str, Any], primary: dict[str, Any]
    ) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
        """Collect every displayed page of one unchanged immutable result."""
        metadata = primary["table_color_metadata"]
        require(metadata["status"] == "ready", "snapshot was not prepared")
        rows = list(primary["data"])
        styles = list(metadata["styles"])
        if payload["form_data"]["server_pagination"]:
            for offset in range(len(rows), metadata["filtered_rowcount"], 20):
                page_payload = self.request_for(
                    payload,
                    metadata["snapshot_id"],
                    metadata["selections"],
                    offset=offset,
                )
                page = self.primary(self.query(page_payload))
                rows.extend(page["data"])
                styles.extend(page["table_color_metadata"]["styles"])
        require(len(rows) == metadata["filtered_rowcount"], "page collection lost rows")
        require(len(styles) == len(rows), "paint and row alignment differs")
        return rows, styles

    def check_zero_new_sql(self, before: int) -> None:
        """Snapshot interactions must not repeat actual ClickHouse SELECTs."""
        if self.api.transport == "test-client":
            require(
                self.counter.count == before,
                "snapshot interaction repeated a source SELECT",
            )

    def native_rows(
        self, payload: dict[str, Any], row_count: int
    ) -> list[dict[str, Any]]:
        """Execute the identical native query to preserve existing calculation rules."""
        records: list[dict[str, Any]] = []
        page_size = 20 if payload["form_data"]["server_pagination"] else row_count
        for offset in range(0, row_count, max(1, page_size)):
            native = copy.deepcopy(payload)
            native["queries"][0].pop("table_color_filter")
            native["queries"][0]["row_offset"] = offset
            result = self.query(native).body["result"][0]
            require(result["status"] == "success", "native comparison query failed")
            records.extend(result["data"])
        require(len(records) == row_count, "native comparison row count differs")
        return records

    def check_raw(self, server: bool) -> None:
        """Verify both pagination modes, all three palettes and cross-column AND."""
        payload = self.make_payload(server=server)
        before_build = self.counter.count
        original = self.primary(self.query(payload))
        after_build = self.counter.count
        if self.api.transport == "test-client":
            require(after_build > before_build, "real SELECT observer did not run")
        baseline, paints = self.complete(payload, original)
        require(len(baseline) == 180, "raw baseline fixture count differs")
        metadata = original["table_color_metadata"]
        require(
            set(metadata["catalog"]["revenue"]) == {"GREEN", "YELLOW", "RED"},
            "raw baseline lacks a palette",
        )
        expected_paints = {
            row["sale_id"]: paint for row, paint in zip(baseline, paints, strict=True)
        }

        def palette(row: dict[str, Any]) -> str:
            """Derive the fixture's expected color independently of metadata."""
            return (
                "RED"
                if row["revenue"] < 0
                else "YELLOW"
                if row["revenue"] <= 100
                else "GREEN"
            )

        selections_to_check = [
            (["GREEN"], False),
            (["GREEN", "YELLOW"], False),
            (["GREEN", "YELLOW", "RED"], False),
            (["GREEN", "YELLOW", "RED"], True),
        ]
        for colors, multiple_columns in selections_to_check:
            selections = [{"column": "revenue", "colors": colors}]
            if multiple_columns:
                selections.extend(
                    [
                        {"column": "region", "colors": ["YELLOW"]},
                        {"column": "quantity", "colors": ["GREEN"]},
                    ]
                )
            expected = [
                row
                for row in baseline
                if palette(row) in colors
                and (
                    not multiple_columns
                    or (row["region"].startswith("A") and row["quantity"] > 4)
                )
            ]
            selected_payload = self.request_for(
                payload, metadata["snapshot_id"], selections
            )
            selected = self.primary(self.query(selected_payload))
            selected_rows, selected_styles = self.complete(selected_payload, selected)
            require(
                selected_rows == expected,
                "selected rows disagree with the independent predicate",
            )
            require(
                selected_styles
                == [expected_paints[row["sale_id"]] for row in selected_rows],
                "filtering changed baseline paint",
            )
        cleared = self.primary(
            self.query(self.request_for(payload, metadata["snapshot_id"], []))
        )
        require(
            cleared["data"] == original["data"],
            "clearing did not restore original first page",
        )
        self.check_zero_new_sql(after_build)
        self.record(
            "RAW_SERVER" if server else "RAW_CLIENT",
            baseline_rows=len(baseline),
            build_selects=after_build - before_build
            if self.api.transport == "test-client"
            else None,
            followup_selects=self.counter.count - after_build
            if self.api.transport == "test-client"
            else None,
        )
        self.check_xlsx(payload, original, baseline, expected_paints)

    @staticmethod
    def cell_bar_theme() -> dict[str, str]:
        """Read configured display tokens without using the color-style resolver."""
        from superset.views.base import (  # pylint: disable=import-outside-toplevel
            get_theme_bootstrap_data,
        )

        tokens = get_theme_bootstrap_data()["theme"]["default"].get("token") or {}
        return {
            key: tokens.get(key, fallback)
            for key, fallback in (
                ("colorSuccess", "#5ac189"),
                ("colorWarning", "#fcc700"),
                ("colorError", "#e04355"),
            )
        }

    @staticmethod
    def expected_default_bars(
        baseline: list[dict[str, Any]], server: bool, theme: dict[str, str]
    ) -> list[dict[str, Any]]:
        """Freeze native signed-bar geometry in each original pagination context."""
        expected: list[dict[str, Any]] = []
        page_size = 20 if server else len(baseline)
        for offset in range(0, len(baseline), page_size):
            page = baseline[offset : offset + page_size]
            maxima = {
                column: max(abs(float(row[column])) for row in page)
                for column in ("sale_id", "quantity", "revenue")
            }
            for row in page:
                paint: dict[str, Any] = {column: {"colors": []} for column in row}
                for column, maximum in maxima.items():
                    if not maximum:
                        continue
                    value = float(row[column])
                    width = abs(math.floor(value / maximum * 100 + 0.5))
                    if width:
                        scheme = "colorError" if value < 0 else "colorSuccess"
                        paint[column]["cellBar"] = {
                            "color": theme[scheme] + "50",
                            "width": width,
                            "offset": 0,
                            "min": 0,
                            "max": maximum,
                        }
                expected.append(paint)
        return expected

    def check_cell_bar_snapshot(
        self,
        payload: dict[str, Any],
        original: dict[str, Any],
        baseline: list[dict[str, Any]],
        expected_paints: list[dict[str, Any]],
        case_id: str,
        export_colors: list[str],
    ) -> None:
        """Select, clear and export frozen paints without indexing default bars."""
        metadata = original["table_color_metadata"]
        expected_catalog = {
            column: [
                color
                for color in ("GREEN", "YELLOW", "RED")
                if any(color in paint[column]["colors"] for paint in expected_paints)
            ]
            for column in baseline[0]
        }
        require(metadata["catalog"] == expected_catalog, "cell-bar catalog differs")
        require(
            metadata["capabilities"]["revenue"] == {"enabled": True, "supported": True},
            "enabled cell-bar target lost its filter capability",
        )
        before = self.counter.count
        selected_counts = []
        for colors in (["GREEN"], ["YELLOW"], ["RED"], ["GREEN", "YELLOW", "RED"], []):
            indices = [
                index
                for index, paint in enumerate(expected_paints)
                if not colors or set(colors) & set(paint["revenue"]["colors"])
            ]
            selections = [{"column": "revenue", "colors": colors}] if colors else []
            selected_payload = self.request_for(
                payload, metadata["snapshot_id"], selections
            )
            selected = self.primary(self.query(selected_payload))
            selected_metadata = selected["table_color_metadata"]
            require(
                selected_metadata["snapshot_id"] == metadata["snapshot_id"]
                and selected_metadata["catalog"] == expected_catalog,
                "cell-bar interaction replaced the baseline snapshot or catalog",
            )
            rows, paints = self.complete(selected_payload, selected)
            require(
                rows == [baseline[index] for index in indices],
                "cell-bar filtering disagrees with explicit formatter matches",
            )
            require(
                paints == [expected_paints[index] for index in indices],
                "cell-bar filtering changed frozen native paint or geometry",
            )
            selected_counts.append(len(rows))
        exported = self.request_for(
            payload,
            metadata["snapshot_id"],
            [{"column": "revenue", "colors": export_colors}] if export_colors else [],
        )
        exported.update(
            result_format="xlsx",
            result_type="results",
            result_format_options={"styled": True},
        )
        expected_rows = [
            row
            for row, paint in zip(baseline, expected_paints, strict=True)
            if not export_colors or set(export_colors) & set(paint["revenue"]["colors"])
        ]
        self.assert_xlsx(
            self.query(exported),
            expected_rows,
            {
                row["sale_id"]: paint
                for row, paint in zip(baseline, expected_paints, strict=True)
            },
        )
        self.check_zero_new_sql(before)
        mode = "SERVER" if payload["form_data"]["server_pagination"] else "CLIENT"
        self.record(
            f"{case_id}_{mode}",
            baseline_rows=len(baseline),
            selected_counts=selected_counts,
            exported_rows=len(expected_rows),
        )

    def check_formatted_cell_bars(
        self,
        server: bool,
        baseline: list[dict[str, Any]],
        default_paints: list[dict[str, Any]],
        theme: dict[str, str],
        *,
        mixed: bool,
    ) -> None:
        """Separate explicit bars from default bars and text/background paint."""
        rules = [
            self.rule(
                "revenue",
                "colorWarning",
                ">",
                100,
                objectFormatting="TEXT_COLOR" if mixed else "CELL_BAR",
            )
        ]
        if mixed:
            rules.append(
                self.rule(
                    "revenue",
                    "colorError",
                    "<",
                    0,
                    objectFormatting="BACKGROUND_COLOR",
                )
            )
        payload = self.make_payload(server=server, show_cell_bars=True, rules=rules)
        original = self.primary(self.query(payload))
        rows, paints = self.complete(payload, original)
        require(rows == baseline, "cell-bar formatting changed source rows")
        expected = copy.deepcopy(default_paints)
        red, green, blue = ImageColor.getcolor(theme["colorWarning"], "RGB")
        for row, row_paint in zip(baseline, expected, strict=True):
            paint = row_paint["revenue"]
            if mixed and row["revenue"] < 0:
                paint.update(backgroundColor=theme["colorError"], colors=["RED"])
                paint.pop("cellBar", None)
            if row["revenue"] > 100:
                if mixed:
                    paint.update(
                        textColor=f"rgb({red}, {green}, {blue})", colors=["YELLOW"]
                    )
                elif "cellBar" in paint:
                    # The native renderer removes the suffix even for solid colors.
                    paint["cellBar"]["color"] = theme["colorWarning"][:-2] + "99"
                    paint["colors"] = ["YELLOW"]
        require(paints == expected, "conditional formatting changed native cell bars")
        self.check_cell_bar_snapshot(
            payload,
            original,
            baseline,
            expected,
            "MIXED_DEFAULT_CELL_BARS" if mixed else "EXPLICIT_CELL_BARS",
            ["YELLOW", "RED"] if mixed else ["YELLOW"],
        )

    def check_cell_bars(self, server: bool) -> None:
        """Keep default bars visible but require matching explicit filter colors."""
        native_payload = self.make_payload(server=server, show_cell_bars=True, rules=[])
        native = self.native_rows(native_payload, 180)
        self.query(native_payload, 400, "TABLE_COLOR_FILTER_INVALID")
        require(
            any(row["revenue"] < 0 for row in native)
            and any(0 < row["revenue"] <= 100 for row in native)
            and any(row["revenue"] > 100 for row in native),
            "cell-bar fixture lacks matched and unmatched signed values",
        )
        payload = self.make_payload(
            server=server,
            show_cell_bars=True,
            rules=[
                self.rule(
                    "revenue", "colorWarning", "is null", objectFormatting="CELL_BAR"
                )
            ],
        )
        original = self.primary(self.query(payload))
        baseline, paints = self.complete(payload, original)
        require(baseline == native, "default bars changed ordinary native rows")
        theme = self.cell_bar_theme()
        expected = self.expected_default_bars(baseline, server, theme)
        require(
            paints == expected,
            "unmatched rule indexed default bars or changed their native painting",
        )
        self.check_cell_bar_snapshot(
            payload, original, baseline, expected, "DEFAULT_ONLY_CELL_BARS", []
        )
        for mixed in (False, True):
            self.check_formatted_cell_bars(
                server, baseline, expected, theme, mixed=mixed
            )

    def check_xlsx(
        self,
        payload: dict[str, Any],
        original: dict[str, Any],
        baseline: list[dict[str, Any]],
        paints: dict[object, dict[str, Any]],
    ) -> None:
        """Read the real XLSX response and compare its rows and static fills."""
        before = self.counter.count
        exported = self.request_for(
            payload,
            original["table_color_metadata"]["snapshot_id"],
            [{"column": "revenue", "colors": ["GREEN", "RED"]}],
        )
        exported["result_format"] = "xlsx"
        exported["result_type"] = "results"
        exported["result_format_options"] = {"styled": True}
        response = self.query(exported)
        expected_indices = [
            index
            for index, row in enumerate(baseline)
            if row["revenue"] < 0 or row["revenue"] > 100
        ]
        expected = [baseline[index] for index in expected_indices]
        self.assert_xlsx(response, expected, paints)
        self.check_zero_new_sql(before)
        mode = "SERVER" if payload["form_data"]["server_pagination"] else "CLIENT"
        self.record(
            f"XLSX_{mode}",
            exported_rows=len(expected),
            workbook_bytes=len(response.content),
        )

        # Current-view references model local search/sort without submitting row data.
        view_rows = expected_indices[-6:][::-1]
        current_view = copy.deepcopy(exported)
        current_view["queries"][0]["table_color_filter"]["view_rows"] = view_rows
        current_response = self.query(current_view)
        self.assert_xlsx(
            current_response, [baseline[index] for index in view_rows], paints
        )
        rejected = copy.deepcopy(current_view)
        rejected["queries"][0]["table_color_filter"]["view_rows"] = [
            view_rows[0],
            view_rows[0],
        ]
        duplicate_response = self.query(rejected, 400)
        require(
            "TABLE_COLOR_FILTER_INVALID" in json.dumps(duplicate_response.body),
            "duplicate current-view references lack a schema validation error",
        )
        unselected = next(
            index for index in range(len(baseline)) if index not in expected_indices
        )
        rejected["queries"][0]["table_color_filter"]["view_rows"] = [unselected]
        self.query(rejected, 400, "TABLE_COLOR_FILTER_INVALID")
        self.check_zero_new_sql(before)
        self.record(
            f"CURRENT_VIEW_XLSX_{mode}",
            exported_rows=len(view_rows),
            invalid_reference_cases=2,
        )

    @staticmethod
    def assert_xlsx(
        response: ApiResponse,
        expected: list[dict[str, Any]],
        paints: dict[object, dict[str, Any]],
    ) -> None:
        """Verify XLSX values and colors against the unchanged authorized snapshot."""
        require(response.content.startswith(b"PK"), "XLSX response is not a workbook")
        require(
            ".xlsx" in response.headers.get("Content-Disposition", ""),
            "XLSX filename is absent",
        )
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.worksheets[0]
        header = [cell.value for cell in sheet[1]]
        data_bars = {
            str(conditional.sqref): rule.dataBar
            for conditional in sheet.conditional_formatting
            for rule in sheet.conditional_formatting[conditional]
            if rule.type == "dataBar"
        }
        require(
            sheet.max_row - 1 == len(expected), "XLSX exported a different row count"
        )
        for cells, expected_row in zip(
            sheet.iter_rows(min_row=2), expected, strict=True
        ):
            actual = {
                str(name): cell.value for name, cell in zip(header, cells, strict=True)
            }
            require(
                actual == expected_row, "XLSX row values differ from selected snapshot"
            )
            for name, cell in zip(header, cells, strict=True):
                ColorAcceptance.assert_xlsx_cell_paint(
                    cell,
                    paints[expected_row["sale_id"]][str(name)],
                    data_bars.get(cell.coordinate),
                )

    @staticmethod
    def assert_xlsx_cell_paint(
        cell: Cell, paint: dict[str, Any], data_bar: DataBar | None
    ) -> None:
        """Check explicit text/fills and both default and explicit frozen bar scales."""
        for css_key, actual in (
            ("backgroundColor", cell.fill.fgColor),
            ("textColor", cell.font.color),
        ):
            if color := paint.get(css_key):
                red, green, blue = ImageColor.getcolor(color, "RGB")
                require(
                    actual is not None
                    and actual.rgb == f"FF{red:02X}{green:02X}{blue:02X}",
                    "XLSX text or fill differs from final native palette",
                )
        if bar := paint.get("cellBar"):
            require(data_bar is not None, "XLSX omitted a visible frozen cell bar")
            assert data_bar is not None
            red, green, blue = ImageColor.getcolor(bar["color"], "RGB")
            require(
                data_bar.color.rgb == f"FF{red:02X}{green:02X}{blue:02X}",
                "XLSX bar color differs from frozen native painting",
            )
            require(len(data_bar.cfvo) == 2, "XLSX bar bounds are absent")
            for bound, key in zip(data_bar.cfvo, ("min", "max"), strict=True):
                require(
                    bound.type == "num" and math.isclose(bound.val, bar[key]),
                    "XLSX recalculated the frozen baseline bar scale",
                )
        else:
            require(data_bar is None, "XLSX added an absent or suppressed cell bar")

    def check_capacity(self, server: bool, rows: int) -> None:
        """Check actual candidate rows around the approved 1,000-result budget."""
        payload = self.make_payload(server=server, rows=rows)
        primary = self.primary(self.query(payload))
        metadata = primary["table_color_metadata"]
        require(
            metadata["status"] == ("ready" if rows <= 1000 else "unavailable"),
            "capacity status differs",
        )
        require(
            len(primary["data"]) == (20 if server else rows),
            "capacity fallback broke ordinary browsing",
        )
        if rows <= 1000:
            require(
                metadata["baseline_rowcount"] == rows, "capacity count was truncated"
            )
        else:
            require(
                "snapshot_id" not in metadata, "oversized result published a snapshot"
            )
            require(
                metadata["reason"]["code"] == "TABLE_COLOR_FILTER_LIMIT_EXCEEDED",
                "wrong capacity reason",
            )
            payload["force"] = False
            payload["queries"][0]["table_color_filter"]["selections"] = [
                {"column": "revenue", "colors": ["GREEN"]}
            ]
            self.query(payload, 422, "TABLE_COLOR_FILTER_LIMIT_EXCEEDED")
        self.record(
            f"CAPACITY_{'SERVER' if server else 'CLIENT'}_{rows}", candidate_rows=rows
        )

    def check_aggregate(self, server: bool, float_percentage: bool = False) -> None:
        """Cover saved/adhoc metrics, percentages and string target composition."""
        percent_metric = "qa_float_revenue" if float_percentage else "gross_revenue"
        percent_column = f"%{percent_metric}"
        rules = [
            self.rule("qa_profit", "colorSuccess", ">", 0),
            self.rule("qa_profit", "colorError", "<", 0),
            self.rule(percent_column, "colorSuccess", ">", 0),
            self.rule(percent_column, "colorError", "<", 0),
            self.rule("region", "colorWarning", "begins with", "A"),
        ]
        payload = self.make_payload(
            server=server,
            rows=20_000,
            aggregate=True,
            percentage=True,
            float_percentage=float_percentage,
            rules=rules,
        )
        original_response = self.query(payload)
        original = self.primary(original_response)
        baseline, paints = self.complete(payload, original)
        metadata = original["table_color_metadata"]
        require(
            0 < len(baseline) <= 1000, "aggregate source exceeds the acceptance scope"
        )
        if float_percentage:
            denominator = original_response.body["result"][-1]["data"][0][
                percent_metric
            ]
            require(
                all(
                    math.isclose(
                        row[percent_column],
                        row[percent_metric] / denominator,
                        rel_tol=1e-9,
                        abs_tol=1e-12,
                    )
                    for row in baseline
                ),
                "Float64 percentage did not use the original totals denominator",
            )
        else:
            # Native contribution excludes Decimal/object totals. Preserve its page
            # context instead of introducing a new percentage definition in a filter.
            require(
                baseline == self.native_rows(payload, len(baseline)),
                "Decimal percentage changed from the native query",
            )
        selections = [
            {"column": "qa_profit", "colors": ["GREEN"]},
            {"column": percent_column, "colors": ["GREEN"]},
            {"column": "region", "colors": ["YELLOW"]},
        ]
        expected_indices = [
            index
            for index, row in enumerate(baseline)
            if row["qa_profit"] > 0
            and row[percent_column] > 0
            and row["region"].startswith("A")
        ]
        require(
            expected_indices and len(expected_indices) < len(baseline),
            "aggregate fixture did not exercise filtering",
        )
        selected_payload = self.request_for(
            payload, metadata["snapshot_id"], selections
        )
        before = self.counter.count
        selected, selected_paints = self.complete(
            selected_payload, self.primary(self.query(selected_payload))
        )
        require(
            selected == [baseline[index] for index in expected_indices],
            "aggregate selection changed rows or percent values",
        )
        require(
            selected_paints == [paints[index] for index in expected_indices],
            "aggregate selection changed color",
        )
        self.check_zero_new_sql(before)
        self.record(
            ("FLOAT_PERCENT_" if float_percentage else "DECIMAL_NATIVE_PERCENT_")
            + ("SERVER" if server else "CLIENT"),
            baseline_groups=len(baseline),
            selected_groups=len(selected),
        )

    def make_time_payload(
        self, server: bool, rules: list[dict[str, Any]]
    ) -> dict[str, Any]:
        """Use a fixture with real datetime metadata without changing existing data."""
        from superset.extensions import db  # pylint: disable=import-outside-toplevel
        from superset.models.slice import (
            Slice,  # pylint: disable=import-outside-toplevel
        )

        source = db.session.query(Slice).filter_by(id=1).one_or_none()
        require(
            source is not None
            and source.viz_type == "table"
            and source.datasource.table_name == "vq_dim_customer",
            "existing customer Table fixture is unavailable",
        )
        require(
            source.datasource.database.url_object.host
            in {"localhost", "127.0.0.1", "::1"},
            "time comparison fixture must be local",
        )
        require(
            any(
                column.column_name == "signup_date" and column.is_dttm
                for column in source.datasource.columns
            ),
            "customer fixture datetime metadata is unavailable",
        )
        self.protected_charts[source.id] = hashlib.sha256(
            (source.params or "").encode()
        ).hexdigest()
        chart_id, datasource_id = source.id, source.datasource_id
        columns = ["country_code", "segment", "is_active"]
        metrics: list[object] = [
            "count",
            {
                "expressionType": "SQL",
                "sqlExpression": "SUM(customer_id)",
                "label": "qa_amount",
                "hasCustomLabel": True,
            },
        ]
        time_range = "2025-06-01 : 2025-07-01"
        form_data = {
            **copy.deepcopy(source.form_data),
            "viz_type": "table",
            "query_mode": "aggregate",
            "slice_id": chart_id,
            "datasource": f"{datasource_id}__table",
            "groupby": columns,
            "all_columns": [],
            "metrics": metrics,
            "percent_metrics": [],
            "server_pagination": server,
            "server_page_length": 20,
            "row_limit": 50_000,
            "show_totals": False,
            "show_cell_bars": False,
            "column_config": {},
            "time_range": time_range,
            "granularity_sqla": "signup_date",
            "time_compare": ["1 month ago"],
            "comparison_type": "values",
            "comparison_color_enabled": False,
            "adhoc_filters": [],
            "conditional_formatting": rules,
            "_color_acceptance_run": self.run_id,
            "_color_acceptance_case": uuid.uuid4().hex,
        }
        source_query = {
            "columns": [],
            "metrics": [
                {
                    "expressionType": "SQL",
                    "sqlExpression": expression,
                    "label": label,
                    "hasCustomLabel": True,
                }
                for label, expression in (
                    ("qa_rows", "COUNT(*)"),
                    (
                        "qa_main_rows",
                        "sumIf(1, signup_date >= toDate('2025-06-01') "
                        "AND signup_date < toDate('2025-07-01'))",
                    ),
                    (
                        "qa_history_rows",
                        "sumIf(1, signup_date >= toDate('2025-05-01') "
                        "AND signup_date < toDate('2025-06-01'))",
                    ),
                )
            ],
            "time_range": "No filter",
            "row_limit": 1,
            "post_processing": [],
        }
        source_response = self.query(
            {
                "datasource": {"id": datasource_id, "type": "table"},
                "queries": [source_query],
                "form_data": {"viz_type": "table", "slice_id": chart_id},
                "result_format": "json",
                "result_type": "full",
                "force": True,
            }
        )
        source_counts = source_response.body["result"][0]["data"][0]
        require(
            source_counts["qa_rows"] >= 10_000
            and source_counts["qa_main_rows"] > 0
            and source_counts["qa_history_rows"] > 0,
            "customer fixture does not cover both comparison periods",
        )
        key = self.create_draft(form_data, datasource_id, chart_id)
        query: dict[str, Any] = {
            "columns": columns,
            "metrics": metrics,
            "granularity": "signup_date",
            "time_range": time_range,
            "time_offsets": ["1 month ago"],
            "orderby": [[column, True] for column in columns],
            "row_limit": 20 if server else 50_000,
            "row_offset": 0,
            "post_processing": [],
        }
        queries = [query]
        if server:
            queries.append(
                {
                    **copy.deepcopy(query),
                    "is_rowcount": True,
                    "time_offsets": [],
                    "row_limit": 50_000,
                }
            )
        query["table_color_filter"] = {
            "version": 2,
            "form_data_key": key,
            "selections": [],
        }
        return {
            "datasource": {"id": datasource_id, "type": "table"},
            "form_data": form_data,
            "queries": queries,
            "force": True,
            "result_type": "full",
            "result_format": "json",
        }

    def check_time_comparison(self, server: bool) -> None:
        """Filter every native comparison display column with one historical query."""
        columns = [
            "Main qa_amount",
            "# qa_amount",
            "△ qa_amount",
            "% qa_amount",
        ]
        rules = [
            self.rule(column, scheme, operator, 0)
            for column in columns
            for scheme, operator in (("colorSuccess", ">"), ("colorError", "<"))
        ]
        payload = self.make_time_payload(server, rules)
        before_build = self.counter.count
        original = self.primary(self.query(payload))
        after_build = self.counter.count
        baseline, paints = self.complete(payload, original)
        metadata = original["table_color_metadata"]
        require(20 < len(baseline) <= 1000, "time comparison lacks multiple pages")
        require(
            all(column in metadata["catalog"] for column in columns),
            "time comparison display columns were not indexed",
        )
        if self.api.transport == "test-client":
            require(
                after_build - before_build == 2,
                "time comparison did not execute exactly one base and one history SQL",
            )
        self.check_zero_new_sql(after_build)
        require(
            baseline == self.native_rows(payload, len(baseline)),
            "time comparison differs from original native result pages",
        )

        def displayed(row: dict[str, Any], column: str) -> float:
            """Independently reproduce the native four comparison display values."""
            main = float(row.get("qa_amount") or 0)
            previous = float(row.get("qa_amount__1 month ago") or 0)
            if column.startswith("Main "):
                return main
            if column.startswith("# "):
                return previous
            difference = main - previous
            if column.startswith("△ "):
                return difference
            if not main and not previous:
                return 0
            if not main or not previous:
                return 1 if main else -1
            return difference / abs(previous)

        before_selections = self.counter.count
        selected_counts = []
        for column in columns:
            color = (
                "GREEN"
                if any(displayed(row, column) > 0 for row in baseline)
                else "RED"
            )
            expected = [
                index
                for index, row in enumerate(baseline)
                if (
                    displayed(row, column) > 0
                    if color == "GREEN"
                    else displayed(row, column) < 0
                )
            ]
            require(expected, "time comparison fixture has no selected display values")
            selected_payload = self.request_for(
                payload,
                metadata["snapshot_id"],
                [{"column": column, "colors": [color]}],
            )
            selected, selected_paints = self.complete(
                selected_payload, self.primary(self.query(selected_payload))
            )
            require(
                selected == [baseline[index] for index in expected],
                "time comparison selection differs from final display color",
            )
            require(
                selected_paints == [paints[index] for index in expected],
                "time comparison selection recomputed original painting",
            )
            selected_counts.append(len(selected))
        self.check_zero_new_sql(before_selections)
        self.record(
            "TIME_COMPARISON_SERVER" if server else "TIME_COMPARISON_CLIENT",
            baseline_groups=len(baseline),
            selected_counts=selected_counts,
            baseline_selects=after_build - before_build
            if self.api.transport == "test-client"
            else None,
            interaction_selects=self.counter.count - before_selections
            if self.api.transport == "test-client"
            else None,
        )

    def check_gradient_and_errors(self) -> None:
        """Verify explicit gradient limits, unrelated targets and API failures."""
        rules = self.raw_rules()
        for rule in rules:
            if rule["column"] == "revenue":
                rule["useGradient"] = True
        payload = self.make_payload(server=False, rules=rules)
        original = self.primary(self.query(payload))
        metadata = original["table_color_metadata"]
        capability = metadata["capabilities"]["revenue"]
        require(
            capability["enabled"] and not capability["supported"],
            "gradient target was not explicitly disabled",
        )
        require(
            capability["reason"]["code"] == "TABLE_COLOR_FILTER_GRADIENT_UNSUPPORTED",
            "gradient reason missing",
        )
        self.query(
            self.request_for(
                payload,
                metadata["snapshot_id"],
                [{"column": "revenue", "colors": ["GREEN"]}],
            ),
            422,
            "TABLE_COLOR_FILTER_GRADIENT_UNSUPPORTED",
        )
        valid = self.primary(
            self.query(
                self.request_for(
                    payload,
                    metadata["snapshot_id"],
                    [{"column": "quantity", "colors": ["GREEN"]}],
                )
            )
        )
        require(
            0 < len(valid["data"]) < len(original["data"]),
            "unrelated pure-color target was disabled by a gradient",
        )
        self.record("GRADIENT_CAPABILITY", unaffected_selected_rows=len(valid["data"]))
        stale = self.request_for(payload, uuid.uuid4().hex, [])
        self.query(stale, 410, "TABLE_COLOR_FILTER_SNAPSHOT_EXPIRED")
        invalid = self.request_for(
            payload,
            metadata["snapshot_id"],
            [{"column": "absent_column", "colors": ["GREEN"]}],
        )
        self.query(invalid, 400, "TABLE_COLOR_FILTER_INVALID")
        foreign = copy.deepcopy(payload)
        foreign["queries"][0]["table_color_filter"]["form_data_key"] = uuid.uuid4().hex
        self.query(foreign, 403, "TABLE_COLOR_FILTER_ACCESS_DENIED")
        anonymous = self.api.request(
            "POST", "/api/v1/chart/data", payload, anonymous=True
        )
        require(
            anonymous.status in {400, 401, 403},
            "unauthenticated chart request succeeded",
        )
        self.record("API_FAILURE_BOUNDARIES", verified_errors=4)

    def check_cross_worker_cache(self) -> None:
        """Read an HTTP-worker snapshot from this distinct Flask app/Redis client."""
        if self.api.transport != "http":
            return
        payload = self.make_payload(server=False)
        original = self.primary(self.query(payload))
        selected_payload = self.request_for(
            payload,
            original["table_color_metadata"]["snapshot_id"],
            [{"column": "revenue", "colors": ["GREEN"]}],
        )
        peer = LocalApi(
            self.api.app, self.api.user_id, self.api.base_url, "test-client"
        )
        before = self.counter.count
        response = peer.request("POST", "/api/v1/chart/data", selected_payload)
        require(response.status == 200, "another worker could not read the snapshot")
        primary = self.primary(response)
        indices = [
            index for index, row in enumerate(original["data"]) if row["revenue"] > 100
        ]
        require(
            primary["data"] == [original["data"][index] for index in indices],
            "cross-worker selected rows differ",
        )
        require(
            primary["table_color_metadata"]["styles"]
            == [original["table_color_metadata"]["styles"][index] for index in indices],
            "cross-worker frozen painting differs",
        )
        require(self.counter.count == before, "another worker repeated source SQL")
        peer.session.close()
        self.record(
            "CROSS_WORKER_REDIS", selected_rows=len(indices), followup_selects=0
        )

    def cleanup(self) -> None:
        """Delete only the temporary drafts created by this process."""
        from superset.extensions import db  # pylint: disable=import-outside-toplevel
        from superset.models.slice import (
            Slice,  # pylint: disable=import-outside-toplevel
        )

        deleted = 0
        for key in self.drafts:
            response = self.api.request("DELETE", f"/api/v1/explore/form_data/{key}")
            require(response.status in {200, 404}, "temporary draft cleanup failed")
            deleted += 1
        for chart_id, fingerprint in self.protected_charts.items():
            persisted = db.session.query(Slice).filter_by(id=chart_id).one()
            persisted_fingerprint = hashlib.sha256(
                (persisted.params or "").encode()
            ).hexdigest()
            require(
                persisted_fingerprint == fingerprint,
                "saved chart parameters changed",
            )
        self.record("CLEANUP", deleted_drafts=deleted, saved_charts_changed=0)

    def run(self, smoke: bool = False, cell_bars_only: bool = False) -> None:
        """Run real cases, preserving cleanup even if an assertion fails."""
        with self.counter.observe(self.chart.engine_spec):
            try:
                if cell_bars_only:
                    for server in (False, True):
                        self.check_cell_bars(server)
                    return
                self.check_raw(False)
                if not smoke:
                    self.check_raw(True)
                    for server in (False, True):
                        self.check_cell_bars(server)
                        for rows in (999, 1000, 1001):
                            self.check_capacity(server, rows)
                        self.check_aggregate(server)
                        self.check_aggregate(server, float_percentage=True)
                        self.check_time_comparison(server)
                    self.check_gradient_and_errors()
                    self.check_cross_worker_cache()
            finally:
                self.cleanup()


def main() -> int:
    """Refuse remote endpoints and run against the existing local test account."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--transport", choices=("http", "test-client"), default="http")
    parser.add_argument("--base-url", default="http://127.0.0.1:8088")
    parser.add_argument("--chart-id", type=int, default=3)
    parser.add_argument("--username", default="admin")
    scope = parser.add_mutually_exclusive_group()
    scope.add_argument("--smoke", action="store_true")
    scope.add_argument("--cell-bars-only", action="store_true")
    args = parser.parse_args()
    endpoint = urlparse(args.base_url)
    require(
        endpoint.scheme in {"http", "https"}
        and endpoint.hostname in {"127.0.0.1", "localhost", "::1"}
        and endpoint.username is None,
        "only loopback service URLs are permitted",
    )
    require(
        os.environ.get("SUPERSET_CONFIG_PATH"),
        "SUPERSET_CONFIG_PATH must reference the existing local config",
    )
    for name in ("NO_PROXY", "no_proxy"):
        entries = os.environ.get(name, "").split(",")
        os.environ[name] = ",".join(
            dict.fromkeys([*entries, "localhost", "127.0.0.1", "::1"])
        )
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
    logging.disable(logging.CRITICAL)
    from superset.app import create_app  # pylint: disable=import-outside-toplevel

    app = create_app()
    from superset.extensions import (  # pylint: disable=import-outside-toplevel
        db,
        security_manager,
    )
    from superset.models.slice import Slice  # pylint: disable=import-outside-toplevel

    started = time.monotonic()
    with app.app_context():
        chart = db.session.query(Slice).filter_by(id=args.chart_id).one_or_none()
        require(
            chart is not None and chart.viz_type == "table",
            "source classic Table chart is absent",
        )
        require(
            chart.datasource.database.backend in {"clickhouse", "clickhousedb"},
            "source is not ClickHouse",
        )
        database_url = chart.datasource.database.url_object
        require(
            database_url.host in {"localhost", "127.0.0.1", "::1"},
            "source database must be local",
        )
        user = security_manager.find_user(username=args.username)
        require(
            user is not None and user.is_active,
            "existing local test account is unavailable",
        )
        source_chart = SourceChart(
            id=chart.id,
            datasource_id=chart.datasource_id,
            form_data=copy.deepcopy(chart.form_data),
            params=chart.params or "",
            query_context=json.loads(chart.query_context or "{}"),
            engine_spec=chart.datasource.database.db_engine_spec,
        )
        api = LocalApi(app, user.id, args.base_url, args.transport)
        suite = ColorAcceptance(api, source_chart)
        suite.run(smoke=args.smoke, cell_bars_only=args.cell_bars_only)
        print(
            json.dumps(
                {
                    "summary": "PASS",
                    "transport": args.transport,
                    "cases": len(suite.outcomes),
                    "api_requests": api.requests,
                    "fixture_selects": suite.counter.count
                    if args.transport == "test-client"
                    else None,
                    "elapsed_seconds": round(time.monotonic() - started, 2),
                }
            ),
            flush=True,
        )
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except AcceptanceError as error:
        print(json.dumps({"summary": "FAIL", "reason": str(error)}), flush=True)
        raise SystemExit(1) from None
    except Exception as error:  # pylint: disable=broad-except
        print(
            json.dumps({"summary": "FAIL", "reason": type(error).__name__}), flush=True
        )
        raise SystemExit(1) from None
