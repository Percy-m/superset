# Licensed to the Apache Software Foundation (ASF) under one
# or more contributor license agreements.  See the NOTICE file
# distributed with this work for additional information
# regarding copyright ownership.  The ASF licenses this file
# to you under the Apache License, Version 2.0 (the
# "License"); you may not use this file except in compliance
# with the License.  You may obtain a copy of the License at
#
#   http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing,
# software distributed under the License is distributed on an
# "AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY
# KIND, either express or implied.  See the License for the
# specific language governing permissions and limitations
# under the License.

import tempfile
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from uuid import UUID, uuid4

import pytest
from flask import g
from openpyxl import load_workbook

from superset import db, security_manager
from superset.commands.dashboard.exceptions import (
    DashboardXlsxInvalidTabError,
    DashboardXlsxNoTableError,
)
from superset.commands.dashboard.export_xlsx import DashboardXlsxExportResult
from superset.daos.dashboard import EmbeddedDashboardDAO
from superset.models.dashboard import Dashboard
from superset.models.embedded_dashboard import EmbeddedDashboard
from superset.models.slice import Slice
from superset.security.guest_token import GuestTokenResourceType
from superset.utils import json
from tests.integration_tests.base_tests import SupersetTestCase
from tests.integration_tests.conftest import with_feature_flags
from tests.integration_tests.fixtures.birth_names_dashboard import (  # noqa: F401
    load_birth_names_dashboard_with_slices,
    load_birth_names_data,
)
from tests.integration_tests.fixtures.query_context import get_query_context


def _clear_request_user() -> None:
    for attribute in ("user", "_login_user"):
        if hasattr(g, attribute):
            delattr(g, attribute)


def _cleanup_guest_export_objects(
    dashboard_id: int,
    baseline_embedded_uuids: set[UUID],
    wrong_dashboard_slug: str,
    guest_role_name: str,
) -> None:
    db.session.rollback()
    for candidate in (
        db.session.query(EmbeddedDashboard).filter_by(dashboard_id=dashboard_id).all()
    ):
        if candidate.uuid not in baseline_embedded_uuids:
            db.session.delete(candidate)
    persisted_dashboard = (
        db.session.query(Dashboard).filter_by(slug=wrong_dashboard_slug).one_or_none()
    )
    if persisted_dashboard is not None:
        db.session.delete(persisted_dashboard)
    if (persisted_role := security_manager.find_role(guest_role_name)) is not None:
        db.session.delete(persisted_role)
    db.session.commit()


class TestDashboardXlsxApi(SupersetTestCase):
    """Exercise HTTP gating, validation, error mapping, and file lifecycle."""

    endpoint = "/api/v1/dashboard/17/export_xlsx/"

    @with_feature_flags(
        STYLED_XLSX_EXPORT=False,
        DASHBOARD_TAB_XLSX_EXPORT=False,
    )
    def test_feature_flags_disabled_return_not_found(self) -> None:
        self.login("admin")

        response = self.client.post(
            self.endpoint,
            json={"tabIds": ["TAB-a"], "dataMask": {}},
        )

        assert response.status_code == 404

    @with_feature_flags(
        STYLED_XLSX_EXPORT=True,
        DASHBOARD_TAB_XLSX_EXPORT=True,
    )
    @patch("superset.dashboards.api.DashboardDAO.get_by_id_or_slug")
    def test_invalid_request_body_returns_bad_request(self, mock_get_dashboard) -> None:
        self.login("admin")
        mock_get_dashboard.return_value = object()

        response = self.client.post(
            self.endpoint,
            json={"tabIds": ["TAB-a", "TAB-a"], "dataMask": {}},
        )

        assert response.status_code == 400
        assert response.json["message"] == {
            "_schema": ["tabIds must not contain duplicate values"]
        }

    @with_feature_flags(
        STYLED_XLSX_EXPORT=True,
        DASHBOARD_TAB_XLSX_EXPORT=True,
    )
    @patch("superset.dashboards.api.ExportDashboardXlsxCommand")
    @patch("superset.dashboards.api.DashboardDAO.get_by_id_or_slug")
    def test_success_returns_xlsx_and_removes_file_on_close(
        self,
        mock_get_dashboard,
        mock_command,
    ) -> None:
        self.login("admin")
        mock_get_dashboard.return_value = object()
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as temporary:
            workbook = Path(temporary.name)
        workbook.write_bytes(b"PK\x03\x04complete")
        mock_command.return_value.run.return_value = DashboardXlsxExportResult(
            workbook,
            "dashboard_17.xlsx",
        )

        with patch(
            "superset.dashboards.api.security_manager.can_access", return_value=True
        ):
            response = self.client.post(
                self.endpoint,
                json={"tabIds": ["TAB-a"], "dataMask": {"3": {}}},
            )

        assert response.status_code == 200
        assert response.data == b"PK\x03\x04complete"
        assert response.headers["Content-Disposition"].endswith(
            "filename=dashboard_17.xlsx"
        )
        response.close()
        assert not workbook.exists()

    @with_feature_flags(
        STYLED_XLSX_EXPORT=True,
        DASHBOARD_TAB_XLSX_EXPORT=True,
    )
    @patch("superset.dashboards.api.ExportDashboardXlsxCommand")
    @patch("superset.dashboards.api.DashboardDAO.get_by_id_or_slug")
    def test_command_errors_map_to_400_and_422(
        self,
        mock_get_dashboard,
        mock_command,
    ) -> None:
        self.login("admin")
        mock_get_dashboard.return_value = object()
        with patch(
            "superset.dashboards.api.security_manager.can_access", return_value=True
        ):
            mock_command.return_value.run.side_effect = DashboardXlsxInvalidTabError()
            invalid_tab = self.client.post(
                self.endpoint,
                json={"tabIds": ["TAB-stale"]},
            )
            mock_command.return_value.run.side_effect = DashboardXlsxNoTableError()
            no_table = self.client.post(
                self.endpoint,
                json={"tabIds": ["TAB-a"]},
            )

        assert invalid_tab.status_code == 400
        assert no_table.status_code == 422

    @pytest.mark.usefixtures("load_birth_names_dashboard_with_slices")
    @with_feature_flags(
        EMBEDDED_SUPERSET=True,
        STYLED_XLSX_EXPORT=True,
        DASHBOARD_TAB_XLSX_EXPORT=True,
    )
    def test_guest_export_uses_trusted_dashboard_context_and_dataset_rls(
        self,
    ) -> None:
        dashboard = db.session.query(Dashboard).filter_by(slug="births").one()
        chart = db.session.query(Slice).filter_by(slice_name="Girls").one()
        dataset = self.get_birth_names_dataset()
        layout = {
            "ROOT_ID": {
                "id": "ROOT_ID",
                "type": "ROOT",
                "children": ["GRID_ID"],
            },
            "GRID_ID": {
                "id": "GRID_ID",
                "type": "GRID",
                "children": ["TAB-GUEST-XLSX"],
                "parents": ["ROOT_ID"],
            },
            "TAB-GUEST-XLSX": {
                "id": "TAB-GUEST-XLSX",
                "type": "TAB",
                "children": ["CHART-GUEST-XLSX"],
                "parents": ["ROOT_ID", "GRID_ID"],
                "meta": {"text": "Guest XLSX"},
            },
            "CHART-GUEST-XLSX": {
                "id": "CHART-GUEST-XLSX",
                "type": "CHART",
                "children": [],
                "parents": ["ROOT_ID", "GRID_ID", "TAB-GUEST-XLSX"],
                "meta": {
                    "chartId": chart.id,
                    "height": 50,
                    "sliceName": chart.slice_name,
                    "width": 12,
                },
            },
        }
        wrong_dashboard_slug = f"guest-xlsx-wrong-{uuid4().hex}"
        guest_role_name = f"guest_xlsx_{uuid4().hex}"
        baseline_embedded_uuids = {item.uuid for item in dashboard.embedded}
        self.addCleanup(
            _cleanup_guest_export_objects,
            dashboard.id,
            baseline_embedded_uuids,
            wrong_dashboard_slug,
            guest_role_name,
        )
        wrong_dashboard = self.insert_dashboard(
            dashboard_title="Guest XLSX wrong dashboard",
            slug=wrong_dashboard_slug,
            owners=[],
            slices=[chart],
            position_json=json.dumps(layout),
            published=True,
        )
        form_data = {
            "dashboardId": wrong_dashboard.id,
            "datasource": f"{dataset.id}__table",
            "groupby": ["gender"],
            "metrics": ["sum__num"],
            "row_limit": 10,
            "time_range": "No filter",
            "viz_type": "table",
        }
        query_context = get_query_context("birth_names", form_data=form_data)
        query = query_context["queries"][0]
        query.update(
            {
                "columns": ["gender"],
                "filters": [],
                "granularity": None,
                "metrics": [{"label": "sum__num"}],
                "orderby": [],
                "row_limit": 10,
                "time_range": "No filter",
            }
        )
        dashboard.position_json = json.dumps(layout)
        chart.params = json.dumps(form_data)
        chart.query_context = json.dumps(query_context)
        embedded = EmbeddedDashboardDAO.upsert(dashboard, [])

        guest_role = security_manager.add_role(guest_role_name)
        expected_permissions = {
            ("can_csv", "Superset"),
            ("can_read", "Dashboard"),
        }
        for permission_name, view_menu_name in expected_permissions:
            permission = security_manager.find_permission_view_menu(
                permission_name,
                view_menu_name,
            )
            assert permission is not None
            security_manager.add_permission_role(guest_role, permission)
        db.session.commit()

        try:
            actual_permissions = {
                (permission.permission.name, permission.view_menu.name)
                for permission in guest_role.permissions
            }
            assert actual_permissions == expected_permissions
            assert not actual_permissions & {
                ("all_database_access", "all_database_access"),
                ("all_datasource_access", "all_datasource_access"),
                ("database_access", dataset.database.perm),
                ("datasource_access", dataset.perm),
                ("schema_access", dataset.schema_perm),
            }

            token = security_manager.create_guest_access_token(
                {"username": "guest_xlsx"},
                [
                    {
                        "type": GuestTokenResourceType.DASHBOARD,
                        "id": str(embedded.uuid),
                    }
                ],
                [{"dataset": str(dataset.id), "clause": "gender = 'girl'"}],
            )
            token_header = token.decode() if isinstance(token, bytes) else token
            headers = {self.app.config["GUEST_TOKEN_HEADER_NAME"]: token_header}
            request_body = {
                "tabIds": ["TAB-GUEST-XLSX"],
                "dataMask": {"dashboardId": wrong_dashboard.id},
            }

            with patch.dict(
                self.app.config,
                {"GUEST_ROLE_NAME": guest_role.name},
            ):
                self.logout()
                _clear_request_user()
                response = self.client.post(
                    f"/api/v1/dashboard/{dashboard.id}/export_xlsx/",
                    json=request_body,
                    headers=headers,
                )

                assert response.status_code == 200
                assert response.content_type.startswith(
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
                payload = response.data
                response.close()
                workbook = load_workbook(
                    BytesIO(payload),
                    read_only=True,
                    data_only=True,
                )
                try:
                    worksheet = workbook[chart.slice_name]
                    rows = worksheet.iter_rows(values_only=True)
                    headers_row = next(rows)
                    gender_index = headers_row.index("gender")
                    genders = {
                        row[gender_index]
                        for row in rows
                        if row[gender_index] is not None
                    }
                    assert genders == {"girl"}
                finally:
                    workbook.close()

                _clear_request_user()
                wrong_response = self.client.post(
                    f"/api/v1/dashboard/{wrong_dashboard.id}/export_xlsx/",
                    json=request_body,
                    headers=headers,
                )

                assert wrong_response.status_code == 404
                assert not wrong_response.content_type.startswith(
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
                wrong_response.close()
        finally:
            _clear_request_user()
            _cleanup_guest_export_objects(
                dashboard.id,
                baseline_embedded_uuids,
                wrong_dashboard_slug,
                guest_role_name,
            )
