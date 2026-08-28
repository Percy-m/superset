# Licensed to the Apache Software Foundation (ASF) under one
# or more contributor license agreements.  See the NOTICE file
# distributed with this work for additional information
# regarding copyright ownership.  The ASF licenses this file
# to you under the Apache License, Version 2.0 (the
# "License"); you may not use this file except in compliance
# with the License.  You may obtain a copy of the License at
#
#   http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing,
# software distributed under the License is distributed on an
# "AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY
# KIND, either express or implied.  See the License for the
# specific language governing permissions and limitations
# under the License.

import time
from decimal import Decimal
from io import BytesIO
from typing import Any
from unittest.mock import Mock, patch
from zipfile import ZipFile

from flask import Flask
from openpyxl import load_workbook

from superset.common.chart_data import ChartDataResultFormat
from superset.common.table_color_query import TableColorQueryProcessor


def _export(rows: list[dict[str, Any]], styles: list[dict[str, Any]]) -> dict[str, Any]:
    """Exercise the real snapshot export boundary without a database connection."""
    processor = object.__new__(TableColorQueryProcessor)
    processor.context = Mock(
        result_format=ChartDataResultFormat.XLSX,
        result_format_options={"styled": True},
        slice_=Mock(slice_name="Precision", id=9),
    )
    processor.trusted = Mock(form_data={})
    processor.selections = [{"column": next(iter(rows[0])), "colors": ["GREEN"]}]
    snapshot = {
        "display_columns": list(rows[0]),
        "display_coltypes": [0] * len(rows[0]),
        "records": rows,
        "snapshot_id": "test-snapshot",
        "generation": "test-generation",
        "source_page_size": 20,
        "catalog": {},
        "capabilities": {},
        "expires_at": time.time() + 60,
    }
    app = Flask(__name__)
    app.config["ROW_LIMIT"] = 1000
    with (
        app.app_context(),
        patch(
            "superset.common.table_color_query.is_feature_enabled", return_value=True
        ),
    ):
        return processor._export(snapshot, rows, styles, None, list(range(len(rows))))[
            0
        ]


def test_snapshot_export_preserves_uint64_next_to_null() -> None:
    """Pandas must not turn nullable UInt64 values into imprecise float64."""
    rows: list[dict[str, Any]] = [{"id": 18446744073709551615}, {"id": None}]
    styles: list[dict[str, Any]] = [{"id": {"colors": []}}, {"id": {"colors": []}}]
    result = _export(rows, styles)
    sheet = load_workbook(BytesIO(result["data"]))["Precision"]
    assert sheet["A2"].value == "18446744073709551615"
    assert sheet["A2"].data_type == "s"
    assert sheet["A3"].value is None


def test_filtered_totals_preserve_decimal_and_uint64_without_numpy_overflow() -> None:
    """Totals sum displayed numeric values, including Decimal-backed columns."""
    rows = [
        {
            "id": 18446744073709551615,
            "amount": Decimal("0.1"),
            "text": "a",
            "bool": True,
        },
        {
            "id": 18446744073709551615,
            "amount": Decimal("0.2"),
            "text": "b",
            "bool": False,
        },
    ]
    totals = TableColorQueryProcessor._sum_display_rows(rows)
    assert totals["id"] == 36893488147419103230
    assert totals["amount"] == Decimal("0.3")
    assert "text" not in totals
    assert "bool" not in totals


def test_filtered_totals_allow_nulls_without_rounding_large_integers() -> None:
    """A NULL does not coerce the rest of a displayed integral column to float."""
    totals = TableColorQueryProcessor._sum_display_rows(
        [{"value": 18446744073709551615}, {"value": None}, {"value": 1}]
    )
    assert totals["value"] == 18446744073709551616
    assert isinstance(totals["value"], int)


def test_filtered_totals_ignore_array_and_json_display_fields() -> None:
    """Non-numeric Table fields must not break filtering during totals creation."""
    totals = TableColorQueryProcessor._sum_display_rows(
        [
            {"tags": [1, 2], "details": {"a": 1}, "value": 10},
            {"tags": [], "details": None, "value": 20},
        ]
    )
    assert totals == {"value": 30}


def test_snapshot_export_keeps_original_gradient_and_bar_range() -> None:
    """Filtering down to one row must not stretch its color or data-bar scale."""
    result = _export(
        [{"amount": 10, "ratio": 0.1}],
        [
            {
                "amount": {
                    "backgroundColor": "#ff000080",
                    "textColor": "rgb(0, 255, 0)",
                    "colors": ["GREEN", "RED"],
                },
                "ratio": {
                    "cellBar": {
                        "color": "#00ff0099",
                        "width": 10,
                        "offset": 0,
                        "min": 0,
                        "max": 1,
                    },
                    "colors": ["GREEN"],
                },
            }
        ],
    )
    workbook = load_workbook(BytesIO(result["data"]))
    assert workbook["Precision"]["A2"].fill.fgColor.rgb == "FFFF7F7F"
    assert workbook["Precision"]["A2"].font.color.rgb == "FF00FF00"
    assert workbook["Precision"]["B2"].value == 0.1
    with ZipFile(BytesIO(result["data"])) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode()
    assert '<cfvo type="num" val="1"/>' in xml
    assert 'rgb="FF00FF00"' in xml
