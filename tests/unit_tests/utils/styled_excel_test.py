# Licensed to the Apache Software Foundation (ASF) under one
# or more contributor license agreements.  See the NOTICE file
# distributed with this work for additional information
# regarding copyright ownership.  The ASF licenses this file
# to you under the Apache License, Version 2.0 (the
# "License"); you may not use this file except in compliance
# with the License.  You may obtain a copy of the License at
#
#   http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing,
# software distributed under the License is distributed on an
# "AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY
# KIND, either express or implied.  See the License for the
# specific language governing permissions and limitations
# under the License.

from datetime import datetime
from io import BytesIO
from zipfile import ZipFile

import pandas as pd
import pytest
from openpyxl import load_workbook

from superset.common.table_alerts import ResolvedTableStyleRule
from superset.utils.core import GenericDataType
from superset.utils.styled_excel import (
    dataframe_to_styled_xlsx,
    resolve_sheet_name,
    StyledExcelError,
    XLSX_DATA_LIMIT_MESSAGE,
    XLSX_STYLE_ALIGNMENT_MESSAGE,
)


def test_styled_excel_preserves_safe_values_and_clickhouse_uint64() -> None:
    dataframe = pd.DataFrame(
        {
            "danger": ["=1+1", "+cmd", "-cmd", "@cmd", "\tcmd", "\rcmd"],
            "long_id": [18446744073709551615] * 6,
            "text": ["汉字😀", "a\x01b", "line\nbreak", "ok", "ok", "ok"],
        }
    )

    payload = dataframe_to_styled_xlsx(
        dataframe,
        [],
        sheet_name="安全导出",
    )
    workbook = load_workbook(BytesIO(payload), data_only=False)
    worksheet = workbook["安全导出"]

    assert [worksheet.cell(row=row, column=1).value for row in range(2, 8)] == [
        "'=1+1",
        "'+cmd",
        "'-cmd",
        "'@cmd",
        "'\tcmd",
        "'_x000D_cmd",
    ]
    assert worksheet["B2"].value == "18446744073709551615"
    assert worksheet["B2"].data_type == "s"
    assert worksheet["C2"].value == "汉字😀"
    assert worksheet["C3"].value == "ab"
    assert worksheet["C4"].value == "line\nbreak"


def test_styled_excel_writes_last_matching_styles_and_native_data_bar() -> None:
    dataframe = pd.DataFrame(
        {
            "gross_revenue": [-100, 50, 200],
            "gross_profit": [10, 20, 30],
        }
    )
    rules = [
        ResolvedTableStyleRule(
            source_column="gross_revenue",
            target_column="gross_revenue",
            dimension="background",
            color="#F5222D",
            operator="<",
            target_value=0,
            use_gradient=False,
        ),
        ResolvedTableStyleRule(
            source_column="gross_revenue",
            target_column="gross_revenue",
            dimension="font",
            color="#FFFFFF",
            operator="<",
            target_value=0,
            use_gradient=False,
        ),
        ResolvedTableStyleRule(
            source_column="gross_profit",
            target_column="gross_profit",
            dimension="data_bar",
            color="#1677FF",
            operator=">",
            target_value=0,
            use_gradient=False,
        ),
    ]

    payload = dataframe_to_styled_xlsx(
        dataframe,
        rules,
        sheet_name="Styled",
    )
    workbook = load_workbook(BytesIO(payload), data_only=False)
    worksheet = workbook["Styled"]

    assert worksheet["A2"].fill.fgColor.rgb == "FFF5222D"
    assert worksheet["A2"].font.color is not None
    assert worksheet["A2"].font.color.rgb == "FFFFFFFF"
    assert worksheet["A3"].fill.patternType is None
    with ZipFile(BytesIO(payload)) as archive:
        sheet_xml = archive.read("xl/worksheets/sheet1.xml")
    assert b"dataBar" in sheet_xml


def test_resolve_sheet_name_is_stable_for_invalid_long_and_duplicate_names() -> None:
    first, first_changed = resolve_sheet_name(
        "Revenue [APAC]: 2026/08?*\\ very long dashboard title",
        7,
        set(),
    )
    second, second_changed = resolve_sheet_name(first, 8, {first})

    assert first == "Revenue APAC 2026 08 very long"
    assert first_changed is True
    assert second == "Revenue APAC 2026 08 very l (2)"
    assert second_changed is True
    assert len(first) <= 31
    assert len(second) <= 31


def test_styled_excel_rejects_a_cell_larger_than_one_mibibyte() -> None:
    dataframe = pd.DataFrame({"value": ["x" * (1024 * 1024 + 1)]})

    with pytest.raises(StyledExcelError, match=XLSX_DATA_LIMIT_MESSAGE):
        dataframe_to_styled_xlsx(dataframe, [], sheet_name="Too large")


def test_styled_excel_applies_temporal_and_saved_number_formats() -> None:
    dataframe = pd.DataFrame(
        {
            "event_time": [datetime(2026, 8, 17, 12, 30)],
            "ratio": [0.125],
        }
    )

    payload = dataframe_to_styled_xlsx(
        dataframe,
        [],
        sheet_name="Formats",
        column_types=[GenericDataType.TEMPORAL, GenericDataType.NUMERIC],
        column_config={"ratio": {"d3NumberFormat": ".1%"}},
    )
    worksheet = load_workbook(BytesIO(payload))["Formats"]

    assert worksheet["A2"].number_format == "yyyy-mm-dd hh:mm:ss"
    assert worksheet["B2"].number_format == "0.0%"


def test_styled_excel_uses_frozen_paints_without_reapplying_rules() -> None:
    """The export keeps source-page paints even if only one matching row remains."""
    dataframe = pd.DataFrame({"profit": [10], "percentage": [0.125]})
    dataframe.attrs["table_color_styles"] = [
        {
            "profit": {
                "backgroundColor": "rgba(150,0,0,0.2)",
                "textColor": "rgb(82, 196, 26)",
                "colors": ["GREEN", "RED"],
            },
            "percentage": {
                "cellBar": {
                    "color": "#faad1499",
                    "width": 25,
                    "offset": 0,
                    "min": 0,
                    "max": 0.5,
                },
                "colors": ["YELLOW"],
            },
        }
    ]
    conflicting_rules = [
        ResolvedTableStyleRule(
            source_column="profit",
            target_column="profit",
            dimension="background",
            color="#1677FF",
            operator="None",
            use_gradient=False,
        )
    ]
    payload = dataframe_to_styled_xlsx(
        dataframe,
        conflicting_rules,
        sheet_name="Frozen",
    )
    worksheet = load_workbook(BytesIO(payload))["Frozen"]
    assert worksheet["A2"].value == 10
    assert worksheet["A2"].fill.fgColor.rgb == "FFEACCCC"
    assert worksheet["A2"].font.color.rgb == "FF52C41A"
    assert worksheet["B2"].value == 0.125
    with ZipFile(BytesIO(payload)) as workbook:
        xml = workbook.read("xl/worksheets/sheet1.xml").decode()
    assert '<cfvo type="num" val="0.5"/>' in xml
    assert 'rgb="FFFAAD14"' in xml
    assert 'negativeBarColorSameAsPositive="1"' in xml


@pytest.mark.parametrize(
    "styles",
    [
        [],
        [{"renamed": {"colors": []}}],
        [{"value": {"colors": ["BLUE"]}}],
        [{"value": {"colors": [], "textColor": None}}],
        [
            {
                "value": {
                    "colors": [],
                    "cellBar": {
                        "color": "#00ff00",
                        "width": 0,
                        "offset": 0,
                        "min": 0,
                        "max": 1,
                    },
                }
            }
        ],
    ],
)
def test_styled_excel_rejects_stale_snapshot_style_alignment(styles: object) -> None:
    """Frame transformations must not silently attach another row/column's paint."""
    dataframe = pd.DataFrame({"value": [1]})
    dataframe.attrs["table_color_styles"] = styles
    with pytest.raises(StyledExcelError, match=XLSX_STYLE_ALIGNMENT_MESSAGE):
        dataframe_to_styled_xlsx(dataframe, [], sheet_name="Stale")


def test_frozen_styles_keep_formula_and_long_integer_safety() -> None:
    """Snapshot styling does not bypass the existing safe value writer."""
    dataframe = pd.DataFrame({"value": [18446744073709551615, "=1+1"]})
    dataframe.attrs["table_color_styles"] = [
        {"value": {"colors": ["GREEN"], "backgroundColor": "#5ac189"}},
        {"value": {"colors": ["RED"], "backgroundColor": "#e04355"}},
    ]
    worksheet = load_workbook(
        BytesIO(dataframe_to_styled_xlsx(dataframe, [], sheet_name="Safe"))
    )["Safe"]
    assert worksheet["A2"].value == "18446744073709551615"
    assert worksheet["A2"].data_type == "s"
    assert worksheet["A3"].value == "'=1+1"


@pytest.mark.parametrize(
    "color,expected",
    [("green", "FF008000"), ("hsl(120, 100%, 50%)", "FF00FF00")],
)
def test_frozen_styles_preserve_named_and_hsl_theme_colors(
    color: str, expected: str
) -> None:
    """A trusted theme need not express semantic base colors as hex strings."""
    dataframe = pd.DataFrame({"value": [1]})
    dataframe.attrs["table_color_styles"] = [
        {"value": {"colors": ["GREEN"], "backgroundColor": color}}
    ]
    sheet = load_workbook(
        BytesIO(dataframe_to_styled_xlsx(dataframe, [], sheet_name="Theme"))
    )["Theme"]
    assert sheet["A2"].fill.fgColor.rgb == expected
